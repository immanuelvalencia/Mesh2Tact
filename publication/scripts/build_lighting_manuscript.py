"""Build the editable methods manuscript using the bundled document runtime."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'publication/manuscript'; OUT.mkdir(parents=True,exist_ok=True)
A=ROOT/'publication/figures/reference_lighting'
m=json.loads((A/'metrics.json').read_text())
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.85)
style=d.styles['Normal']; style.font.name='Times New Roman'; style.font.size=Pt(11)
style.paragraph_format.space_after=Pt(7); style.paragraph_format.line_spacing=1.08
for name in ('Title','Heading 1','Heading 2'):
    d.styles[name].font.name='Times New Roman'; d.styles[name].font.color.rgb=RGBColor(0,0,0)
d.styles['Title'].font.size=Pt(22)
d.styles['Heading 1'].font.size=Pt(15)
d.styles['Caption'].font.color.rgb=RGBColor(0,0,0)
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer=sec.footer.paragraphs[0]; footer.alignment=2
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,1)
def page(): d.add_page_break()
def eq(t):
    para=d.add_paragraph(); math=OxmlElement('m:oMath'); run=OxmlElement('m:r'); text=OxmlElement('m:t'); text.text=t; run.append(text); math.append(run); para._p.append(math)
def fig(file,caption):
    d.add_picture(str(A/file),width=Inches(6.7))
    para=d.add_paragraph(caption); para.style='Caption'; para.paragraph_format.space_after=Pt(10)

d.add_paragraph('Reference guided fitting of edge array illumination for geometric tactile image synthesis','Title')
p('Methods Results and Discussion manuscript draft')
h('Abstract')
p('We describe a reproducible appearance-fitting procedure for a geometric visuotactile simulator using user-supplied GelSight images. The simulator converts an object surface into an indentation height field and computes reflection from colored arrays spanning three sensor edges. A two-layer procedural background is fitted to the pixelwise median of nine spatially separated pyramid-contact frames. Effective RGB array weights are then estimated with bounded linear least squares from one sphere-contact frame using an assumed spherical geometry. The resulting configuration is installed as the default desktop appearance. This study evaluates reconstruction of the fitting images and provides qualitative depth-response examples; it does not estimate force or establish physical calibration.')
p(f'The interior background mean absolute error was {m["background_interior_mae"]:.5f} on a normalized 0–1 RGB scale. Within the selected sphere region, the fitted rendering had mean absolute error {m["sphere_roi_mae"]:.5f}, compared with {m["background_only_roi_mae"]:.5f} for the procedural background alone. These are in-sample errors, not independent generalization estimates. The approximation reproduces broad color organization but lacks the real membrane texture, smooth peripheral deformation, and camera-boundary appearance. The results support use as an adjustable geometric baseline and motivate calibrated geometry, sequence-level evaluation, and explicit optical transport modeling.')
p('Keywords: visuotactile simulation; GelSight; edge illumination; inverse rendering; geometric contact; appearance fitting')
h('Scope of the evidence')
p('This document describes the implemented fit and its observed results at the time of preparation. The target journal has not been specified, so the document uses a neutral single-column manuscript layout. Journal readiness depends on additional experimental validation, dataset provenance, and the selected journal’s author requirements. No claim of Q1 acceptance, novelty priority, photorealism, or sim-to-real transfer is made.')
h('Optical motivation')
p('GelSight observes a reflective membrane through a transparent elastomer. Illumination from different directions produces a color response related to membrane orientation [1]. Example-based tactile simulation has also used contact geometry to predict pixel intensity [2]. The present implementation uses an analytic edge-array approximation with a small image-based fit, rather than reproducing an established sensor calibration procedure.')

page(); h('Methods and source data')
p('The supplied images represent sphere, square-edge, truncated-cone, and pyramid contacts. The numerical fit uses only nine pyramid images and one sphere image. The pyramid inputs are sequence_220 through sequence_228, frame_0000_raw.png in each sequence. The sphere input is hemisphere/video/sequence_003/sequence_003_frame_0012_raw.png. The square-edge and truncated-cone images informed qualitative inspection only; they were not used in the least-squares objective and do not constitute a held-out quantitative test.')
p('All inputs reside under Touchlab-VTS/dataset/GelSight and are read without modifying the original dataset. Images are converted to RGB and resized with Pillow to 246 × 185 pixels. These are encoded RGB values normalized by 255; no radiometric linearization, white-balance correction, lens calibration, or camera response inversion is performed. Consequently, the fitted coefficients are effective image-space coefficients, not physical spectral emission measurements.')
h('Background estimation')
p('A pixelwise median over the nine pyramid images estimates the empty-gel appearance because the visible contacts occur at different spatial locations. This is a robust heuristic; it is not a separately recorded no-contact measurement. Shared contact influence, fixed blemishes, and persistent deformation can remain in the estimate. A 15-pixel border is excluded from the background fitting objective, leaving a 216 × 155 pixel interior region.')
p('The procedural background contains a constant RGB base, radial darkening, and two elliptical Gaussian color layers. The layer centers and spreads are fixed at (0.50, 0.26, 0.40, 0.32) and (0.50, 0.56, 0.25, 0.25), respectively, in normalized image coordinates. Twelve quantities are optimized: base RGB, darkening strength, two layer RGB triplets, and two blending strengths. Each is constrained to [0, 1]. SciPy least_squares uses an initial parameter vector [0.2, 0.3, 0.3, 0.25, 0.3, 0.5, 0.65, 0.5, 0.65, 0.4, 0.5, 0.5] and at most 80 function evaluations.')
eq('B₀ = b;   Bₖ = (1 − αₖ) Bₖ₋₁ + αₖ cₖ;   B = clip[(1 − v r²) B₂]')
p('Here b is the base color, c is a layer color, α is its Gaussian spatial opacity, v is the corner-darkening parameter, and r² is the renderer’s normalized squared radial coordinate. The fit minimizes the sum of squared RGB differences from the median image over the interior mask. The algorithm fits appearance; the Gaussian layers are not a reconstruction of the internal optical field.')
h('Assumed sphere geometry')
p('The sphere center is manually approximated as (480, 350) pixels in the original 984 × 739 frame, with an apparent contact radius of 175 pixels. The physical sensor extent is assumed to be 18.6 × 14.3 mm and sphere radius 4.5 mm. These values are not supplied metrology. A spherical cap is constructed from these assumptions and smoothed with a Gaussian of standard deviation 0.8 pixels. The fitting region is a circle with radius 1.12 times the estimated contact radius.')

page(); h('Edge array rendering and coefficient fitting')
p('The membrane height is the negative indentation depth. Normals are calculated from finite differences with the sensor pixel spacing. Each selected edge is represented by 12 equally spaced midpoint emitters spanning its full length. Their contributions are averaged so array intensity does not grow with the numerical sample count. A source-to-pixel direction and distance are recomputed at every deformed surface point; moving a contact therefore changes its illumination even when its local shape remains the same.')
eq('n = normalize(−∂h/∂x, −∂h/∂y, 1)')
eq('S = a + Σⱼ Σₖ [wⱼ / 12] (d₀ⱼₖ / dⱼₖ)ᶠ max(0, n · lⱼₖ)')
p('S denotes the RGB shading response, a the ambient color, w an effective RGB array coefficient, and l the normalized direction toward an emitter. Distance falloff exponent f is fixed at 0.4 for this fit; edge distance is 1.1 times the corresponding sensor half-size. Elevation is fixed at 25 degrees. Specular gain is zero and diffuse gain is one. The source geometry is an effective model; propagation through acrylic and elastomer, refraction, internal reflection, visibility, scattering, and spectral camera sensitivity are not solved.')
p('Three white-light basis responses are evaluated at bottom, left, and top edges, subtracting each flat-membrane response from its deformed response. The differential responses are attenuated by the procedural overlay’s transmission. For each output channel, three nonnegative weights are fitted independently with scipy.optimize.lsq_linear, bounded from zero to three. The target is the real sphere image minus the median background, restricted to the circular fitting region.')
eq('w꜀ = arg min ‖X꜀ w꜀ − (I꜀ − M꜀)‖²;   0 ≤ w꜀ ≤ 3')
p('X is the stack of differential basis responses after overlay attenuation; I is the sphere image and M the median image. Each fitted row is represented in the application by its maximum coefficient as intensity and its normalized triplet as RGB color. The fit occurs before final clipping, whereas rendering clips optical RGB before overlay composition. The objectives are therefore not exactly identical, and reported final-image errors are computed from the actual rendered image rather than inferred from the optimizer cost.')
p('The direction labels refer to source placement in this renderer’s image coordinates. The fitted red-dominant source is below the image, which produces a red band above the horizontal indentation. It must not be interpreted as a measured physical LED placement. The top source contains substantial red as well as blue, indicating coefficient coupling and an effective appearance fit rather than independent monochromatic source recovery.')
h('Deployment')
p('The full result is stored in configs/sensors/Real GelSight reference.json. On desktop startup without an explicit sensor file, this configuration is loaded and selected. A lighting-only button remains available. The stored effect settings include photon-noise strength 0.002, gel texture strength 0.008, and blur 0.5 pixels. These effects and the application softness setting of 0.03 mm were manually selected, not estimated by the lighting optimizer. The figures below show the deterministic optical fit without those camera effects.')

page(); h('Results of the appearance fit')
fig('comparison.png','Figure 1. Real sphere frame used for fitting (left), approximate simulated sphere (center), and fitted empty-gel appearance (right). All panels are shown at the fitting resolution. The middle panel uses assumed spherical geometry and the fitted optical response without added image effects. This is an in-sample comparison; the missing texture, sharp contact boundary, and omitted camera housing remain visible.')
p(f'The median-background interior MAE was {m["background_interior_mae"]:.5f}, equivalent to {255*m["background_interior_mae"]:.2f} intensity levels on an 8-bit scale. For the {m["sphere_roi_pixels"]:,} pixels in the sphere fitting region, final RGB MAE was {m["sphere_roi_mae"]:.5f} and RMSE was {m["sphere_roi_rmse"]:.5f}. The background-only MAE in the same region was {m["background_only_roi_mae"]:.5f}. These values depend on the hand-selected center, radius, mask, and assumed physical scale. They do not isolate illumination error from geometry or registration error.')
p('The approximate contact reproduces a pink upper region and a blue lower region. The contact is smoother in texture but sharper in outline than the real image. A low global background error can conceal substantial contact mismatch because most pixels are not deformed; the contact-region result is therefore reported separately. No independent sequence-level error, confidence interval, or significance test is justified by this single fitted contact.')
h('Effective array coefficients')
for name,row in zip(('Bottom','Left','Top'),m['effective_weights']):
    p(f'{name} array RGB weights: {row[0]:.6f}, {row[1]:.6f}, {row[2]:.6f}.')
p('The fitted weights show a red-dominant bottom array, a green-dominant left array, and a mixed top array. Their amplitudes should be interpreted jointly with diffuse gain, source geometry, background overlay, and the assumed sphere depth. Multiple parameter combinations can yield similar images; source identifiability has not been established.')
g=m['gel_parameters']
p('Fitted base RGB: '+', '.join(f'{v:.5f}' for v in g['background'])+f'; corner darkening: {g["vignette"]:.5f}. '+
  ' '.join(f'Layer {i+1} RGB '+', '.join(f'{v:.5f}' for v in layer['color'])+f', opacity {layer["strength"]:.5f}.' for i,layer in enumerate(g['glows']))+
  f' The assumed peak sphere depth after smoothing is {m["assumed_peak_depth_mm"]:.3f} mm.')

page(); h('Depth response and qualitative interpretation')
fig('depth_response.png','Figure 2. Approximate sphere contacts (upper row) and horizontal-edge contacts (lower row) rendered at assumed indentation depths of 0.1, 0.4, and 0.8 mm. These are synthetic sensitivity examples, not paired real measurements. The line endpoints are artifacts of the finite synthetic edge support. Increasing depth changes contact extent and shading, but the displayed millimeter values have not been validated against the supplied acquisition sequences.')
p('The synthetic sweep demonstrates that the configuration remains coupled to geometry: increasing indentation enlarges the spherical contact and separates the opposing edge bands. The response is not a simple global brightness multiplier. For the piecewise-linear edge, local slope is largely constant while depth changes the illuminated support. Thus, brightness need not increase monotonically with depth at every pixel. Distance, slope, clipping, and overlay opacity jointly determine the result.')
p('The real square-edge examples show a narrow red band above a blue band; a targeted regression test confirms the same relative channel ordering in the fitted configuration. This is a directional consistency check, not a quantitative match of band width, color saturation, or deformation spread. The supplied truncated-cone and pyramid images reveal position-dependent responses and nonlocal transitions that remain useful future validation cases. Their full geometries, poses, and indentation distances have not been recovered.')
h('Numerical verification')
p('The line-array integrator was compared with a 1024-sample reference integration in a simplified flat-surface case with specularity and falloff disabled. The test requires maximum intensity error below 0.002 and distinguishes the array result from a center-point source by more than 0.02. Additional tests check opposite-edge symmetry, normal sensitivity, and depth sensitivity. These tests establish internal numerical behavior, not agreement with real optical transport. Source images remain unchanged.')

page(); h('Discussion and limitations')
p('The procedure supplies a reproducible intermediate baseline between hand-selected lighting and a sensor-specific calibrated model. Its practical benefit is a small, editable configuration driven by real images while retaining geometric control over object pose and cut depth. The observed agreement is strongest at the level of broad color organization. It is insufficient to claim physical equivalence or photorealistic tactile sensing.')
p('The primary uncertainty is geometry. Sphere radius, contact center, contact extent, sensor size, and indentation are assumed or visually estimated, and they couple directly to recovered lighting strengths. No force, displacement, timing, loading history, or camera calibration accompanies the fitting inputs. The real elastomer may deform beyond the contact boundary, whereas geometric slicing with Gaussian smoothing does not solve elasticity. More accurate lighting alone cannot remove this geometry mismatch.')
p('The background estimator combines nine contact frames rather than a measured no-contact frame. The two Gaussian layers omit the green border, dark housing, sensor blemishes, and fine texture. Encoded RGB fitting also absorbs unknown exposure, channel sensitivity, and gamma into the coefficients. Red energy in the upper array is consistent with this lack of identifiability and should not be presented as hardware characterization. Parameter bounds and fixed source elevation constrain the solution but do not prove uniqueness.')
p('The fitting sphere must not also be treated as an independent test sample. Adjacent frames share object, sensor state, and acquisition conditions; randomly splitting them would overstate generalization. A journal evaluation should separate whole acquisition sequences, fit on controlled contacts, and report held-out errors for spheres, edges, cones, and pyramid tips at known positions and depths. No such evaluation is claimed here.')
h('Experiments needed for submission')
p('Record a true no-contact reference and known-radius sphere indentations at measured depths across a spatial grid, with fixed camera exposure and documented LED arrangement. Hold out entire acquisition trials and at least one contact geometry. Compare the present model against uniform directional lights, edge point sources, full arrays without fitting, and an established calibrated tactile renderer under the same geometry and masks. Report contact-region RGB error, color difference under a documented color pipeline, edge-band width, spatial generalization, temporal consistency, and runtime at preview and capture resolutions. Use trial-level confidence intervals and sensitivity analyses for radius, depth, background estimation, and emitter count.')
h('Reproducibility and references')
p('The fitting script is publication/scripts/fit_reference_lighting.py. Run conda run -n mesh2tact python publication/scripts/fit_reference_lighting.py from the project root. It writes the JSON configuration, metrics.json, and two comparison figures under publication/figures/reference_lighting. Metrics are generated from the actual rendering. The original dataset must remain at the documented sibling path. Software versions and source-data hashes should be archived with the submission; they are not recorded by this prototype script.')
p('[1] Yuan W, Dong S, Adelson EH. GelSight High-Resolution Robot Tactile Sensors for Estimating Geometry and Force. Sensors. 2017;17:2762. https://doi.org/10.3390/s17122762')
p('[2] Si Z, Yuan W. Taxim An Example-based Simulation Model for GelSight Tactile Sensors. 2021 preprint. https://arxiv.org/abs/2109.04027. Cited for related example-based simulation, not as a quantitatively evaluated baseline.')
d.save(OUT/'GelSight_reference_fitting_manuscript.docx')
print(OUT/'GelSight_reference_fitting_manuscript.docx')
